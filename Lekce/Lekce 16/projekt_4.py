import pandas as pd 
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.chart import BarChart, Reference

df = pd.read_csv("test_data_new.csv")

ok_count = df["Status"].value_counts().get("OK", 0)
failed_count = df["Status"].value_counts().get("Failed", 0)# hodnoty ze kterych chceme graf tvorit 

wb = load_workbook("Report_highlighted.xlsx")
ws = wb.active
#listy pro jednotliva data 
ws.append([])
ws.append(["Shrnuti vysledku"])
ws.append(["Pocet OK testu", ok_count])
ws.append(["Shrnuti Failed testu", failed_count])


chart = BarChart()# zakladni struktura
chart.title = "Vysledky testu"
chart.x_axis.title = "Status"
chart.y_axis.title = "Pocet"

data = Reference(ws, min_col=2, min_row=10, max_row=11)
categories = Reference(ws, min_col=2, min_row=10, max_row=11)

chart.add_data(data, titles_from_data=False)
chart.set_categories(categories)

ws.add_chart(chart, "E17")


ws.column_dimensions["B"].width = 50
ws.column_dimensions["D"].width = 50
ws.column_dimensions["E"].width = 50

wb.save("Report_final.xlsx")
